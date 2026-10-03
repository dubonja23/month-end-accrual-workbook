"""JE backup workbook (xlsx)."""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .money import ZERO

MONEY = '#,##0.00;(#,##0.00)'
HEAD_FILL = PatternFill("solid", fgColor="E4EDE9")


def _cell(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (list, dict, set)):
        return str(v)
    return v


def _sheet(wb, title, headers, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([_cell(v) for v in r])
    for idx, h in enumerate(headers, 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = max(10, min(42, len(str(h)) + 4))
        if h in money_cols:
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
    ws.freeze_panes = "A2"
    return ws


def build_workbook(result: dict) -> bytes:
    je = result["je"]
    h = je["header"]
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, "Header", ["Field", "Value"], [
        ["Company", result["company"]], ["Period", result["period_label"]], ["Memo", h["memo"]],
        ["JE date", h["date"].isoformat()], ["Auto-reverse", h["reverse_date"].isoformat()],
        ["Journal", h["journal"]], ["Reference", h["reference"]],
        ["Vendor OpEx", result["totals"]["vendor_opex"]], ["Legal", result["totals"]["legal"]],
        ["Capex Projects", result["totals"]["capex"]], ["Total accrual", result["totals"]["total"]],
        ["Capex final?", "Not final (unallocated GL)" if result["capex"]["not_final"] else "Final"],
        ["Generated", result["generated_at"]],
    ])

    _sheet(wb, "Tie-out", ["Check", "JE", "Expected", "Difference", "Result"],
           [[c["name"], c["actual"], c["expected"], c["difference"], "OK" if c["ok"] else "FAIL"] for c in je["checks"]],
           money_cols=("JE", "Expected", "Difference"))

    prior_by_acct: dict = {}
    for p in result["prior_je"]["lines"]:
        if p["amount"] > 0:
            prior_by_acct[p["acct"]] = prior_by_acct.get(p["acct"], ZERO) + p["amount"]
    cur_by_acct: dict = {}
    names: dict = {}
    for ln in je["lines"]:
        if ln["side"] == "debit":
            cur_by_acct[ln["acct"]] = cur_by_acct.get(ln["acct"], ZERO) + ln["debit"]
            names[ln["acct"]] = ln["acct_name"]
    accts = sorted(set(prior_by_acct) | set(cur_by_acct))
    _sheet(wb, "By GL account", ["Account", "Name", "Current", "Prior month", "Change"],
           [[a, names.get(a, ""), cur_by_acct.get(a, ZERO), prior_by_acct.get(a, ZERO),
             cur_by_acct.get(a, ZERO) - prior_by_acct.get(a, ZERO)] for a in accts],
           money_cols=("Current", "Prior month", "Change"))

    prior_by_key = {f"{p['vendor_id']}|{p['acct']}": p["amount"] for p in result["prior_je"]["lines"]}
    drivers = {r["key"]: r["driver"] for r in result["mom"]["rows"]}
    je_by_key: dict = {}
    for ln in je["lines"]:
        if ln["side"] == "debit" and ln["section"] == "Vendor OpEx":
            je_by_key[ln["ref"]] = je_by_key.get(ln["ref"], ZERO) + ln["debit"]
    opex_headers = ["Vendor", "Vendor ID", "Account", "Type", "Owner", "Annual budget", "Budget source", "Budget YTD",
                    "GL total YTD", "Spend YTD", "Prior-year service (excl)", "Prepaid (incl)", "Billed after cutoff (incl)",
                    "Prepaid next-year (excl)", "Liability opening", "Prior-year bills", "Liability remaining",
                    "Calculated accrual", "Accrual rule", "Standing rule", "Standing basis", "Standing amount",
                    "Override", "Override explanation", "Override by", "Final accrual", "JE debit", "Prior accrual",
                    "Change", "Driver", "Dims source", "Dept", "Location", "Item", "Prior formula"]
    rows = []
    for ln in result["vendor_opex"]["lines"]:
        st, ov = ln.get("standing") or {}, ln.get("override") or {}
        prior = prior_by_key.get(ln["key"], ZERO)
        rows.append([ln["vendor_name"], ln["vendor_id"], ln["gl_account"], ln["vendor_type"], ln["owner"],
                     ln["annual_budget"], ln["budget_source"], ln["budget_ytd"], ln["gl_total_ytd"], ln["spend_ytd"],
                     ln["prior_year_service"], ln["prepaid_included"], ln["after_cutoff"], ln["prepaid_next_year"],
                     ln["liability_opening"], ln["liability_prior_year_bills"], ln["liability_remaining"],
                     ln["formula_accrual"], ln["accrual_rule"], st.get("rule"), st.get("basis"), st.get("amount"),
                     ov.get("amount"), ov.get("explanation"), ov.get("user"), ln["final_accrual"],
                     je_by_key.get(ln["key"], ZERO), prior, ln["final_accrual"] - prior, drivers.get(ln["key"]),
                     ln["dims"]["source"], ln["dims"]["dept"], ln["dims"]["location"], ln["dims"]["item"],
                     ln["prior_formula"]])
    _sheet(wb, "Vendor OpEx detail", opex_headers, rows, money_cols={
        h for h in opex_headers if any(w in h for w in ("budget", "Budget", "GL total", "Spend", "service", "Prepaid",
                                                         "Billed", "Liability", "bills", "accrual", "amount",
                                                         "Override", "JE debit", "Change", "formula"))
        and h not in ("Override explanation", "Override by", "Budget source", "Accrual rule")})

    lrows = []
    for r in result["legal"]["rows"]:
        for i, jl in enumerate(r["je_lines"], 1):
            lrows.append([r["vendor"], r["vendor_id"], i if r["splits"] else "", jl["amount"], jl["gl"], jl["item"],
                          r["customer"], r["project"], r["dept"], r["location"], jl.get("note") or r["note"]])
    _sheet(wb, "Legal detail", ["Vendor", "Vendor ID", "Split line", "Amount", "GL", "Item", "Customer", "Project",
                                "Dept", "Location", "Note"], lrows, money_cols=("Amount",))

    crows = []
    for g in result["capex"]["groups"]:
        crows.append([g["deal"], g["project_id"], g["vendor"], g["vendor_id"], g["location_id"], len(g["line_nos"]),
                      g["cost"], g["net"], g["je_amount"], "Invoiced beyond earned" if g["invoiced_beyond_earned"] else ""])
    _sheet(wb, "Capex detail", ["Deal", "Project", "Vendor", "Vendor ID", "Location", "Lines", "Cost",
                                "Earned less invoiced", "JE amount", "Note"], crows,
           money_cols=("Cost", "Earned less invoiced", "JE amount"))

    _sheet(wb, "JE lines", ["Line", "Section", "Account", "Account desc", "Dept", "Location", "Project", "Customer",
                            "Vendor", "Item", "Memo", "Debit", "Credit", "Dims source"],
           [[ln["line_no"], ln["section"], ln["acct"], ln["acct_name"], ln["dept"], ln["location"], ln["project"],
             ln["customer"], ln["vendor_id"], ln["item"], ln["memo"], ln["debit"], ln["credit"], ln["dims_source"]]
            for ln in je["lines"]], money_cols=("Debit", "Credit"))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
