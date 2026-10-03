"""End-to-end on the generated sample data for 2026-09, plus the CLI exit code."""
import io
import json

import pytest
import yaml
from openpyxl import load_workbook

from accrual_workbook.cli import main
from accrual_workbook.engine.backup import build_workbook
from accrual_workbook.engine.journal import debit_credit_pairs_match
from accrual_workbook.runner import DEFAULT_CONFIG, Workbook
from helpers import D


@pytest.fixture(scope="module")
def result(sample_dir, tmp_path_factory):
    wb = Workbook(sample_dir, DEFAULT_CONFIG, tmp_path_factory.mktemp("db") / "wb.db")
    return wb.run("2026-09")


def line(result, key):
    return next(ln for ln in result["vendor_opex"]["lines"] if ln["key"] == key)


def test_period_totals(result):
    assert result["totals"] == {"vendor_opex": D("253091.51"), "legal": D("88550.00"), "capex": D("61950.00"),
                                "total": D("403591.51")}


def test_je_balanced_and_all_checks_pass(result):
    je = result["je"]
    assert je["total_debits"] == je["total_credits"] == D("403591.51")
    assert len(je["lines"]) == 86
    assert all(c["ok"] for c in je["checks"])
    assert debit_credit_pairs_match(je["lines"])


def test_scope_exclusion_counts(result):
    assert result["scope"]["counts"] == {"removed document": 1, "after cutoff": 22, "non-qualifying journal": 11,
                                         "blank vendor": 2, "out-of-scope account": 12, "legal": 24}


@pytest.mark.parametrize("key,field,expected", [
    ("V1001|60200", "budget_ytd", "135000.00"),      # annual-vs-YTD regression case
    ("V1001|60200", "final_accrual", "0.00"),
    ("V1015|61100", "final_accrual", "5000.00"),     # two roster lines, matched by account
    ("V1015|65300", "final_accrual", "3000.00"),
    ("V1002|63100", "prepaid_included", "18400.00"),  # 184/365 of 36,500
    ("V1002|63100", "prepaid_next_year", "18100.00"),
    ("V1002|63100", "final_accrual", "3600.00"),
    ("V1005|61200", "final_accrual", "12500.00"),    # account rollup
    ("V1022|64100", "final_accrual", "8500.00"),     # vendor rollup
    ("V1027|62300", "liability_remaining", "7000.00"),
    ("V1027|62300", "final_accrual", "17000.00"),
    ("V1032|60100", "final_accrual", "19000.00"),    # roster liability 4,000
    ("V1035|65200", "formula_accrual", "2000.00"),   # accrual rule -> 0
    ("V1035|65200", "final_accrual", "0.00"),
    ("V1003|62100", "final_accrual", "10000.02"),    # dept split line
    ("V1023|62100", "budget_ytd", "33750.00"),       # budget override 45,000
    ("V1031|64100", "budget_ytd", "7500.00"),        # excluded budget line -> roster spread
    ("V1014|61300", "final_accrual", "0.00"),        # zero (billed in full)
    ("V1013|62100", "final_accrual", "7600.00"),     # unbilled 2 x 3,800
    ("V1033|62300", "final_accrual", "3000.00"),     # unbilled 1.5 x 2,000
    ("V1038|62100", "final_accrual", "750.00"),      # no dated bill -> 1 month
    ("V1028|65200", "final_accrual", "977.78"),      # avg 12,400/9 less 400
    ("V1021|64200", "final_accrual", "0.00"),        # rate unless billed (billed)
    ("V1018|63200", "final_accrual", "6000.00"),     # rate unless billed (not billed)
    ("V1004|60100", "final_accrual", "44000.00"),    # budget rule
    ("V1010|61100", "final_accrual", "0.00"),        # generic paid-through, formula 45,000
    ("V1025|65300", "after_cutoff", "4000.00"),
    ("V1040|65100", "final_accrual", "1800.00"),     # carried manual add
])
def test_seeded_lines(result, key, field, expected):
    assert line(result, key)[field] == D(expected)


def test_seeded_flags(result):
    flags = {t["key"]: [f["code"] for f in t["flags"]] for t in result["standing"]["table"]}
    assert flags["V1014|61300"] == ["new_bill_on_zero"]
    assert flags["V1036|62100"] == ["term_ending"]
    assert flags["V1004|60100"] == ["budget_change"]
    assert flags["V1012|64100"] == ["spend_change"]
    assert flags["V1010|61100"] == ["paid_through_removal"]
    assert "no_dated_bill" in flags["V1038|62100"]
    assert flags["V1040|65100"] == ["new_bill"]


def test_capex_and_attention(result):
    assert result["capex"]["not_final"] is True
    assert [g["project_id"] for g in result["capex"]["over_invoiced"]] == ["P-1002"]
    keys = {a["key"] for a in result["attention_candidates"]}
    assert {"dims|V1026|60100", "capex_over|P-1002|V3002", "capex_unallocated"} <= keys


def test_mom_drivers(result):
    drivers = {r["key"]: r["driver"] for r in result["mom"]["rows"]}
    assert drivers["V1019|64300"] == "prior accrual not rebooked"
    assert drivers["V1030|64200"] == "accrual now $0"
    assert drivers["V1040|65100"] == "manual add"
    assert drivers["V2003|61500"] == "new accrual"


def test_backup_workbook_sheets(result):
    wb = load_workbook(io.BytesIO(build_workbook(result)))
    assert wb.sheetnames == ["Header", "Tie-out", "By GL account", "Vendor OpEx detail", "Legal detail",
                             "Capex detail", "JE lines"]


def test_cli_run_writes_outputs_and_exits_zero(sample_dir, tmp_path):
    out = tmp_path / "out"
    code = main(["run", "--period", "2026-09", "--data", str(sample_dir), "--db", str(tmp_path / "w.db"),
                 "--out", str(out)])
    assert code == 0
    assert (out / "je_upload_202609.csv").exists() and (out / "je_backup_202609.xlsx").exists()
    summary = json.loads((out / "summary_202609.json").read_text(encoding="utf-8"))
    assert summary["totals"]["total"] == 403591.51


def test_cli_run_fails_when_a_check_fails(sample_dir, tmp_path):
    cfg = yaml.safe_load(DEFAULT_CONFIG.read_text(encoding="utf-8"))
    cfg["je"]["dept_splits"][0]["shares"][0]["share"] = 0.40  # shares now add to 90%
    bad = tmp_path / "bad.yaml"
    bad.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    code = main(["run", "--period", "2026-09", "--data", str(sample_dir), "--config", str(bad),
                 "--db", str(tmp_path / "w.db"), "--out", str(tmp_path / "out")])
    assert code == 1
